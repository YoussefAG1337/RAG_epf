"""Serve ingested documents to the viewer: the original file and a line-numbered view."""

from __future__ import annotations

import re
import unicodedata
from io import BytesIO
from pathlib import Path
from typing import Literal
from uuid import UUID

from pydantic import BaseModel
from sqlalchemy.orm import Session, sessionmaker

from app.ingestion import IngestionError, resolve_selected_file
from app.models import Document
from app.text_parsing import _cell_text

DocumentKind = Literal["pdf", "markdown", "text", "spreadsheet"]
_MEDIA_TYPES: dict[DocumentKind, str] = {
    "pdf": "application/pdf",
    "markdown": "text/markdown; charset=utf-8",
    "text": "text/plain; charset=utf-8",
    "spreadsheet": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class DocumentMissingError(LookupError):
    """The document is not ingested, or its file is no longer in the source directory."""


class DocumentInfo(BaseModel):
    id: str
    subject: str
    course_id: str
    title: str
    source_filename: str
    kind: DocumentKind
    page_count: int


class TextLine(BaseModel):
    line: int
    text: str


class SheetRow(BaseModel):
    row: int
    cells: list[str]


class Sheet(BaseModel):
    number: int
    name: str
    rows: list[SheetRow]


class DocumentContent(BaseModel):
    """Non-PDF documents as the viewer renders them: numbered lines or sheet rows."""

    kind: DocumentKind
    lines: list[TextLine] = []
    sheets: list[Sheet] = []


def document_kind(source_filename: str) -> DocumentKind:
    suffix = Path(source_filename).suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".xlsx":
        return "spreadsheet"
    if suffix in {".md", ".markdown"}:
        return "markdown"
    return "text"


def media_type(kind: DocumentKind) -> str:
    return _MEDIA_TYPES[kind]


def get_document(session_factory: sessionmaker[Session], document_id: UUID) -> DocumentInfo:
    with session_factory() as session:
        document = session.get(Document, document_id)
        if document is None:
            raise DocumentMissingError("document not found")
        return DocumentInfo(
            id=str(document.id),
            subject=document.subject,
            course_id=document.course_id,
            title=document.title or document.source_filename,
            source_filename=document.source_filename,
            kind=document_kind(document.source_filename),
            page_count=document.page_count,
        )


def document_path(source_directory: Path, info: DocumentInfo) -> Path:
    """The document's file beneath the source directory, never outside it."""

    try:
        path, _ = resolve_selected_file(source_directory, info.source_filename)
    except IngestionError as error:
        raise DocumentMissingError(
            "the document's file is no longer in the course folder; re-run ingestion"
        ) from error
    return path


def document_content(path: Path, kind: DocumentKind) -> DocumentContent:
    """Line-numbered text (matching the ingestion line numbers) or spreadsheet rows."""

    data = path.read_bytes()
    if kind == "spreadsheet":
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            sheets: list[Sheet] = []
            for number, sheet in enumerate(workbook.worksheets, start=1):
                rows: list[SheetRow] = []
                first_row = sheet.min_row or 1
                for row_number, row in enumerate(
                    sheet.iter_rows(values_only=True), start=first_row
                ):
                    cells = [_cell_text(value) for value in row]
                    while cells and not cells[-1]:
                        cells.pop()
                    if cells:
                        rows.append(SheetRow(row=row_number, cells=cells))
                sheets.append(Sheet(number=number, name=sheet.title, rows=rows))
        finally:
            workbook.close()
        return DocumentContent(kind=kind, sheets=sheets)
    if kind == "pdf":
        return DocumentContent(kind=kind)
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    text = unicodedata.normalize("NFC", text.replace("\r\n", "\n").replace("\r", "\n"))
    return DocumentContent(
        kind=kind,
        lines=[
            TextLine(line=number, text=re.sub(r"\s+$", "", line))
            for number, line in enumerate(text.split("\n"), start=1)
        ],
    )
